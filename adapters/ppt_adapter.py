# -*- coding: utf-8 -*-
"""adapters/ppt_adapter.py - PowerPoint (.pptx, .ppt) 변환 어댑터.

- .pptx: python-pptx 로 슬라이드/표/도형 텍스트 네이티브 추출
- .ppt: PowerPoint COM 자동화로 슬라이드 구조 추출
"""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

from adapters.base import BaseAdapter, ConversionResult
from core.text_cleaner import rows_to_table, tidy

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


class PptAdapter(BaseAdapter):
    """PowerPoint 슬라이드 문서 변환기."""

    def can_handle(self, ext: str) -> bool:
        return ext.lower() in {".pptx", ".ppt"}

    def _table_to_content(self, rows: list[list[str]]) -> str:
        """슬라이드 내 표를 메모장에 붙여넣은 것처럼 서식 없이 순수 텍스트로 변환합니다."""
        if not rows:
            return ""

        grid = []
        for r in rows:
            cleaned = [c.strip().replace("\r\n", " ").replace("\r", " ").replace("\n", " ") for c in r]
            if any(cleaned):
                grid.append(cleaned)

        if not grid:
            return ""

        num_cols = len(grid[0])
        lines = []

        # 2열 비교/정의형 표: "• 항목: 설명" 형태로 가독성 극대화
        if num_cols == 2:
            for r in grid:
                col0 = r[0].strip()
                col1 = r[1].strip() if len(r) > 1 else ""
                if col0 and col1:
                    lines.append(f"• {col0}: {col1}")
                elif col0:
                    lines.append(f"• {col0}")
                elif col1:
                    lines.append(f"  {col1}")
            return "\n".join(lines)

        # 다열 표: 메모장 붙여넣기 스타일로 각 셀을 공백과 구분자로 연결
        for r_idx, r in enumerate(grid):
            non_empty = [c for c in r if c]
            if non_empty:
                lines.append("  |  ".join(non_empty))

        return "\n".join(lines)

    def _convert_pptx_native(self, path: Path) -> tuple[str, str, list[tuple[int, str]]]:
        """python-pptx 를 활용하여 Office 없이 슬라이드 내용을 콘텐츠 중심으로 추출합니다."""
        from pptx import Presentation

        prs = Presentation(str(path.resolve()))
        body_parts = []
        pages = []

        for i, slide in enumerate(prs.slides, start=1):
            slide_items = []

            # 슬라이드 내 도형들을 상단 -> 하단, 좌측 -> 우측 순서로 정렬
            sorted_shapes = sorted(
                list(slide.shapes),
                key=lambda s: (getattr(s, "top", 0), getattr(s, "left", 0)),
            )

            for shape in sorted_shapes:
                if shape.has_table:
                    tbl = shape.table
                    rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
                    t_text = self._table_to_content(rows)
                    if t_text:
                        slide_items.append(t_text)
                elif shape.has_text_frame:
                    tf_text = shape.text_frame.text.strip()
                    if tf_text:
                        slide_items.append(tf_text)

            slide_md = "\n\n".join(slide_items)
            pages.append((i, slide_md))
            body_parts.append(f"### 슬라이드 {i}\n\n{slide_md}")

        return tidy("\n\n".join(body_parts)), "", pages

    def _convert_via_com(self, path: Path, session: Any) -> tuple[str, str, list[tuple[int, str]]]:
        """PowerPoint COM 자동화를 활용해 슬라이드 내용을 콘텐츠 중심으로 추출합니다."""
        app = session.get("ppt")
        abs_path = str(path.resolve())
        pres = app.Presentations.Open(abs_path, ReadOnly=-1, Untitled=-1, WithWindow=-1)
        body_parts = []
        pages = []

        try:
            for i in range(1, pres.Slides.Count + 1):
                slide = pres.Slides(i)
                slide_items = []
                for j in range(1, slide.Shapes.Count + 1):
                    shape = slide.Shapes(j)
                    # 표 확인
                    try:
                        if shape.HasTable:
                            tbl = shape.Table
                            rows = [
                                [
                                    tbl.Cell(r, c).Shape.TextFrame.TextRange.Text.strip()
                                    for c in range(1, tbl.Columns.Count + 1)
                                ]
                                for r in range(1, tbl.Rows.Count + 1)
                            ]
                            t_text = self._table_to_content(rows)
                            if t_text:
                                slide_items.append(t_text)
                            continue
                    except Exception:
                        pass

                    # 일반 텍스트 상자
                    try:
                        if shape.HasTextFrame and shape.TextFrame.HasText:
                            txt = shape.TextFrame.TextRange.Text.replace("\r", "\n").strip()
                            if txt:
                                slide_items.append(txt)
                    except Exception:
                        continue

                slide_md = "\n\n".join(slide_items)
                pages.append((i, slide_md))
                body_parts.append(f"### 슬라이드 {i}\n\n{slide_md}")
        finally:
            pres.Close()

        return tidy("\n\n".join(body_parts)), "", pages

    def convert(self, path: Path, session: Any = None) -> ConversionResult:
        ext = path.suffix.lower()
        method = "native"

        if ext == ".pptx":
            try:
                body, tables, pages = self._convert_pptx_native(path)
            except Exception:
                if session is not None:
                    body, tables, pages = self._convert_via_com(path, session)
                    method = "converted"
                else:
                    raise
        else:
            # .ppt
            if session is None:
                raise RuntimeError(".ppt 변환을 위해서는 PowerPoint COM 세션이 필요합니다.")
            body, tables, pages = self._convert_via_com(path, session)
            method = "converted"

        if not body.strip():
            raise RuntimeError("슬라이드에서 텍스트 내용을 추출하지 못했습니다.")

        return ConversionResult(
            body=body,
            tables=tables,
            pages=pages,
            conversion_method=method,
            status="success",
        )

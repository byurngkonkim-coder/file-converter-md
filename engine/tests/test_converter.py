# -*- coding: utf-8 -*-
"""tests/test_converter.py - 통합 문서 변환기 단위 및 E2E 무결성 검증 테스트 스위트."""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
from pathlib import Path

# Python 콘솔 UTF-8 재설정
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# 경로 등록
APP_DIR = Path(__file__).resolve().parent.parent
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

import yaml
from core.batch_runner import convert_batch, convert_one_file, unique_target_path
from core.frontmatter_builder import (
    build_frontmatter,
    generate_doc_id,
    split_frontmatter,
)
from core.markdown_formatter import format_markdown_document
from core.metadata_extractor import (
    detect_category,
    detect_language,
    extract_description,
    extract_keywords_and_topics,
    extract_title,
)
from core.text_cleaner import detect_encoding, rows_to_table, tidy
from core.text_postprocessor import (
    fix_hyphenated_words,
    fix_punctuation_spacing,
    join_broken_paragraphs,
    postprocess_markdown,
)


def test_text_postprocessor():
    print("[1-1] text_postprocessor 가독성 정제 모듈 검증 중...")
    # 1. 문장 중간 하드 줄바꿈 결합 검증
    broken_text = (
        "한국무역정보통신은\n"
        "전자무역 인프라 혁신을\n"
        "지속적으로 추진하고 있습니다.\n\n"
        "새로운 문단입니다."
    )
    joined = join_broken_paragraphs(broken_text)
    assert "한국무역정보통신은 전자무역 인프라 혁신을 지속적으로 추진하고 있습니다." in joined
    assert "새로운 문단입니다." in joined
    print("  - 단락 내 하드 줄바꿈 결합: 통과")

    # 2. 하이픈 어절 결합 검증
    hyphen_text = "스마트 정보통-\n신망 인프라"
    fixed_hyphen = fix_hyphenated_words(hyphen_text)
    assert "스마트 정보통신망 인프라" in fixed_hyphen
    print("  - 하이픈 분절 어절 복원: 통과")

    # 3. 구두점 및 비가시 문자 정리 검증
    punct_text = "이것은 테스트 입니다 .\ufeff 다음 문장 ( 괄호 내용 ) 입니다 ."
    fixed_punct = fix_punctuation_spacing(punct_text)
    assert "입니다." in fixed_punct
    assert "(괄호 내용)" in fixed_punct
    print("  - 구두점 및 괄호 공백 정규화: 통과")

    # 4. 마크다운 보호형 전체 후처리 검증 (코드블록/표 보존)
    md_sample = (
        "---\ntitle: \"테스트\"\n---\n\n"
        "# 제목\n\n"
        "문장이 중간에\n"
        "끊어져 있는 상태입니다.\n\n"
        "```python\n"
        "# 코드 블록은 줄바꿈 보존\n"
        "def func():\n"
        "    return 1\n"
        "```\n\n"
        "| 열 1 | 열 2 |\n"
        "|---|---|\n"
        "| 값 A | 값 B |\n"
    )
    processed_md = postprocess_markdown(md_sample)
    assert "문장이 중간에 끊어져 있는 상태입니다." in processed_md
    assert "def func():\n    return 1" in processed_md  # 코드 블록 보존
    assert "| 열 1 | 열 2 |" in processed_md  # 표 보존
    print("  - 마크다운 구조 보존형 전체 후처리: 통과")


def test_text_cleaner():
    print("[1] text_cleaner 모듈 검증 중...")
    # rows_to_table
    rows = [["헤더1", "헤더2"], ["데이터A", "데이터B|파이프"]]
    tbl_md = rows_to_table(rows)
    assert "| 헤더1 | 헤더2 |" in tbl_md
    assert "\\|파이프" in tbl_md

    # 서술문 표의 메모장 스타일 평탄화 검증
    narrative_rows = [["사례 제목", "이 문장은 표 셀 안에 들어있는 40자 이상의 긴 본문 서술형 문단 내용입니다."]]
    txt_res = rows_to_table(narrative_rows)
    assert "|" not in txt_res
    assert "긴 본문 서술형 문단 내용입니다." in txt_res
    print("  - rows_to_table (정형 표 & 서술문 평탄화): 통과")

    # tidy
    raw = "줄1\n\n\n\n줄2  \n```python\ncode\n\n\ncode\n```\n"
    res = tidy(raw)
    assert "\n\n\n" not in res.split("```")[0]
    assert "code\n\n\ncode" in res
    print("  - tidy: 통과")


def test_frontmatter_builder():
    print("[2] frontmatter_builder 모듈 검증 중...")
    doc_id = generate_doc_id("test/path.docx")
    fm = build_frontmatter(
        doc_id=doc_id,
        title="테스트 문서",
        description="테스트 문서 요약입니다.",
        source_file="sample.docx",
        source_path="D:/docs/sample.docx",
        source_type="docx",
        doc_category="보고서",
        topics=["테스트", "보고서"],
        keywords=["키워드1", "키워드2"],
        entities=["한국무역정보통신"],
        status="success",
    )
    assert fm.startswith("---")
    assert fm.rstrip().endswith("---")
    assert 'type: "document"' in fm
    assert f'id: "{doc_id}"' in fm

    # YAML 파싱 검증
    parsed, _ = split_frontmatter(fm + "\n\n본문 내용")
    assert parsed is not None
    assert parsed["title"] == "테스트 문서"
    assert parsed["keywords"] == ["키워드1", "키워드2"]
    assert parsed["entities"] == ["한국무역정보통신"]
    print("  - build_frontmatter 및 split_frontmatter: 통과")


def test_metadata_extractor():
    print("[3] metadata_extractor 모듈 검증 중...")
    sample_text = (
        "# 2026년 3분기 디지털 무역 동향 보고서\n\n"
        "한국무역정보통신(KTNET)은 전자무역 인프라 혁신을 추진하고 있습니다. "
        "본 보고서는 무역 자동화 및 관세 행정 혁신 성과를 분석합니다."
    )
    title = extract_title(sample_text, "fallback.docx")
    assert "2026년 3분기 디지털 무역 동향 보고서" in title

    cat = detect_category(sample_text, "report.docx")
    assert cat == "보고서"

    lang = detect_language(sample_text)
    assert lang == "ko"

    kws, topics = extract_keywords_and_topics(sample_text, limit=3)
    assert len(kws) > 0

    desc = extract_description(sample_text)
    assert "한국무역정보통신" in desc
    print("  - metadata_extractor: 통과")


def test_markdown_formatter():
    print("[4] markdown_formatter 모듈 검증 중...")
    doc = format_markdown_document(
        title="테스트 문서",
        source_file="test.docx",
        source_path="D:/test.docx",
        doc_category="보고서",
        topics=["무역", "자동화"],
        description="요약 설명",
        body_content="이것은 본문입니다.",
        tables_content="| A | B |\n|---|---|\n| 1 | 2 |",
    )
    assert "# 테스트 문서" in doc
    assert "## 문서 개요" in doc
    assert "## 본문 내용" in doc
    assert "## 표 및 데이터" in doc
    assert "## 추출 품질 및 검토 사항" in doc
    assert "## 원본 정보" in doc
    print("  - format_markdown_document: 통과")


def test_e2e_conversions():
    print("[5] 파일 포맷별 E2E 변환 검증 중...")
    tmp_dir = Path(tempfile.mkdtemp(prefix="test_conv_"))
    out_dir = tmp_dir / "output"

    try:
        # 1. TXT 파일 생성 및 변환
        txt_file = tmp_dir / "01_텍스트_샘플.txt"
        txt_file.write_text(
            "자유와 지혜에 관한 명언\n\n자유는 책임을 동반한다. 배움은 끝이 없는 여정이다.",
            encoding="utf-8",
        )

        # 2. MD 파일 생성 및 변환 (기존 Front Matter 포함)
        md_file = tmp_dir / "02_마크다운_샘플.md"
        md_file.write_text(
            "---\ntitle: 기존 제목\nauthor: 홍길동\n---\n\n# 수정된 제목\n\n기존 마크다운 문서의 내용입니다.",
            encoding="utf-8",
        )

        # 3. DOCX 파일 생성 (python-docx 활용)
        import docx
        doc = docx.Document()
        doc.add_heading("2026년도 사업 계획서", level=1)
        doc.add_paragraph("본 계획서는 2026년도 시스템 구축 방향을 설명합니다.")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "구분"
        table.rows[0].cells[1].text = "예산"
        table.rows[1].cells[0].text = "클라우드"
        table.rows[1].cells[1].text = "1000만원"
        docx_file = tmp_dir / "03_워드_샘플.docx"
        doc.save(str(docx_file))

        # 4. XLSX 파일 생성 (openpyxl 활용)
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "실적현황"
        ws.append(["부서", "목표", "달성률"])
        ws.append(["디지털사업팀", "100", "120%"])
        xlsx_file = tmp_dir / "04_엑셀_샘플.xlsx"
        wb.save(str(xlsx_file))

        # 5. PPTX 파일 생성 (python-pptx 활용)
        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "인공지능 비전 발표"
        slide.placeholders[1].text = "차세대 AI 파이프라인 소개"
        pptx_file = tmp_dir / "05_파워포인트_샘플.pptx"
        prs.save(str(pptx_file))

        # 6. PDF 파일 생성 (PyMuPDF fitz 활용)
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        pdoc = fitz.open()
        page = pdoc.new_page()
        page.insert_text((50, 72), "PDF 테스트 문서 제목", fontsize=18)
        page.insert_text((50, 100), "이 문서는 PyMuPDF 텍스트 레이어로 생성된 PDF 파일입니다.", fontsize=11)
        pdf_file = tmp_dir / "06_PDF_샘플.pdf"
        pdoc.save(str(pdf_file))
        pdoc.close()

        # 7. 이미지 파일 생성 및 OCR 변환 (PIL 활용)
        from PIL import Image, ImageDraw
        img = Image.new('RGB', (400, 100), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        draw.text((20, 35), 'Test OCR Document 2026', fill=(0, 0, 0))
        img_file = tmp_dir / "07_이미지_샘플.png"
        img.save(str(img_file))

        # 배치 일괄 변환 실행 (결과 MD와 리포트 폴더 분리)
        report_dir = tmp_dir / "reports"
        files = [txt_file, md_file, docx_file, xlsx_file, pptx_file, pdf_file, img_file]
        report = convert_batch(files, out_dir=out_dir, report_dir=report_dir, overwrite=True)

        assert report["counts"]["success"] == len(files)
        # 1. 리포트 폴더에 보고서 파일 존재 확인
        assert (report_dir / "_conversion_report.json").exists()
        assert (report_dir / "_conversion_summary.md").exists()
        # 2. 결과_MD 폴더에는 리포트 파일이 없고 순수 마크다운만 존재하는지 확인
        assert not (out_dir / "_conversion_report.json").exists()
        assert not (out_dir / "_conversion_summary.md").exists()

        # 결과 Markdown 파일 개별 확인
        for f in files:
            md_res = out_dir / f"{f.stem}.md"
            assert md_res.exists(), f"결과 파일 누락: {md_res.name}"
            content = md_res.read_text(encoding="utf-8")
            fm, body = split_frontmatter(content)
            assert fm is not None, f"YAML Front Matter 누락: {md_res.name}"
            assert fm["status"] == "success"
            assert "## 문서 개요" in body
            assert "## 본문 내용" in body
            print(f"  - [{f.suffix}] {f.name} -> {md_res.name}: 변환 성공 및 스키마 검증 통과")

    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_markitdown_engine():
    print("[6] Microsoft MarkItDown 통합 엔진 검증 중...")
    from adapters.markitdown_adapter import MarkItDownAdapter, is_markitdown_available
    assert is_markitdown_available(), "markitdown 패키지가 사용 불가능 상태입니다."

    adapter = MarkItDownAdapter()
    tmp_dir = Path(tempfile.mkdtemp(prefix="test_md_engine_"))
    try:
        # 테스트용 워드 문서 생성
        import docx
        doc = docx.Document()
        doc.add_heading("MarkItDown 연동 테스트", level=1)
        doc.add_paragraph("이 문서는 Microsoft MarkItDown 엔진을 통해 변환되는 테스트 문서입니다.")
        docx_file = tmp_dir / "markitdown_test.docx"
        doc.save(str(docx_file))

        # 1. 어댑터 직접 변환 검증
        res = adapter.convert(docx_file)
        assert res.status == "success", f"MarkItDown 어댑터 변환 실패: {res.errors}"
        assert "MarkItDown 연동 테스트" in res.body
        print("  - MarkItDownAdapter 직접 변환: 통과")

        # 2. convert_batch(..., engine='markitdown') 하이브리드 파이프라인 검증
        out_dir = tmp_dir / "output"
        report_dir = tmp_dir / "reports"
        report = convert_batch(
            [docx_file],
            out_dir=out_dir,
            report_dir=report_dir,
            engine="markitdown",
            overwrite=True,
        )
        assert report["counts"]["success"] == 1
        res_file = out_dir / "markitdown_test.md"
        assert res_file.exists()
        content = res_file.read_text(encoding="utf-8")
        fm, body = split_frontmatter(content)
        assert fm is not None
        assert fm["status"] == "success"
        assert "MarkItDown 연동 테스트" in body
        print("  - convert_batch(engine='markitdown') 파이프라인 연동: 통과")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_skip_scanned_option():
    print("[7] 텍스트 레이어 없는 스캔 PDF 및 이미지 제외(skip_scanned) 옵션 검증 중...")
    from PIL import Image
    tmp_dir = Path(tempfile.mkdtemp(prefix="test_skip_scan_"))
    out_dir = tmp_dir / "output"
    report_dir = tmp_dir / "reports"

    try:
        # 1. 일반 텍스트 문서
        txt_file = tmp_dir / "normal.txt"
        txt_file.write_text("일반 텍스트 문서 내용입니다.", encoding="utf-8")

        # 2. 이미지 문서 (OCR 대상)
        img_file = tmp_dir / "scan.png"
        Image.new("RGB", (100, 50), color=(255, 255, 255)).save(str(img_file))

        # 3. 텍스트 레이어 없는 스캔형 PDF (글자 수 0)
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        pdoc = fitz.open()
        pdoc.new_page()
        pdf_file = tmp_dir / "scan_empty.pdf"
        pdoc.save(str(pdf_file))
        pdoc.close()

        # skip_scanned=True 실행
        report = convert_batch(
            [txt_file, img_file, pdf_file],
            out_dir=out_dir,
            report_dir=report_dir,
            skip_scanned=True,
            overwrite=True,
        )

        assert report["counts"]["success"] == 1
        assert report["counts"]["skip"] == 2
        assert report["counts"]["failed"] == 0

        # 결과 확인: normal.md는 생성되고, 스캔 문서 md는 생성되지 않아야 함
        assert (out_dir / "normal.md").exists()
        assert not (out_dir / "scan.md").exists()
        assert not (out_dir / "scan_empty.md").exists()
        print("  - skip_scanned 옵션 적용 시 스캔 문서 제외 및 일반 문서 정상 변환: 통과")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_portable_deployment_and_path_safety():
    print("[8] USB 배포 및 경로 이식성(Portable Path Safety) 검증 중...")
    import unicodedata
    from core.batch_runner import is_same_source, normalize_path
    from core.config import get_ocr_python_candidates, get_safe_temp_dir

    # 1. 드라이브 문자(D: -> E:) 변경 시 동일 출처 인식 검증
    fm_sample = {
        "title": "테스트",
        "source_file": "2026_업무보고.docx",
        "source_path": "D:/usb_drive/docs/2026_업무보고.docx",
    }
    src_on_e_drive = Path("E:/usb_drive/docs/2026_업무보고.docx")
    assert is_same_source(fm_sample, src_on_e_drive), "드라이브 문자 변경 시 출처 매칭 실패"
    print("  - USB 드라이브 문자(D: -> E:) 변경 시 기변환 파일 매칭: 통과")

    # 2. 상대 경로 (source_rel_path) 인식 검증
    fm_with_rel = {
        "title": "테스트",
        "source_file": "guide.docx",
        "source_path": "Z:/some/path/guide.docx",
        "source_rel_path": "manuals/guide.docx",
    }
    src_rel_match = APP_DIR / "manuals" / "guide.docx"
    assert is_same_source(fm_with_rel, src_rel_match), "상대 경로 출처 매칭 실패"
    print("  - 상대 경로(source_rel_path) 기반 출처 매칭: 통과")

    # 3. 맥 OS NFD 자모 분리 -> 윈도우 NFC 정규화 검증
    nfd_name = unicodedata.normalize("NFD", "한글문서.docx")
    normalized_p = normalize_path(f'"{nfd_name}"')
    assert unicodedata.is_normalized("NFC", normalized_p.name), "유니코드 NFC 정규화 실패"
    print("  - 유니코드 NFD/NFC 및 따옴표 경로 정규화: 통과")

    # 4. 안전 임시 폴더 및 Python 후보 목록 무결성 검증
    safe_temp = get_safe_temp_dir()
    assert safe_temp.exists() and safe_temp.is_dir(), "안전 임시 폴더 유효성 실패"
    py_cands = get_ocr_python_candidates()
    assert len(py_cands) > 0, "Python 후보 목록 비어있음"
    assert Path(sys.executable) in py_cands, "현재 인터프리터가 후보에 누락됨"
    print("  - 안전 임시 폴더 및 포터블 Python 후보 동적 탐색: 통과")


def run_all_tests() -> int:
    try:
        test_text_postprocessor()
        test_text_cleaner()
        test_frontmatter_builder()
        test_metadata_extractor()
        test_markdown_formatter()
        test_e2e_conversions()
        test_markitdown_engine()
        test_skip_scanned_option()
        test_portable_deployment_and_path_safety()
        print("\n[성공] 모든 셀프 테스트 및 기능 검증을 성공적으로 통과했습니다! (ALL PASSED)")
        return 0
    except Exception as e:
        import traceback
        print(f"\n[실패] 테스트 실행 중 오류 발생: {e}")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(run_all_tests())
